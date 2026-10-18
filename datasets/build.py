"""Build the evaluation set: 36 messy files from six invented sources, their gold tables and 10 drift scenarios.

Every file is rendered from a clean list of records (the gold table) through the quirks a real export has: odd
delimiters and encodings, preamble and footer lines, merged Excel headers, nested JSON, fixed-width layouts,
decimal commas, currency symbols, split debit/credit columns, Fahrenheit, hours as h:mm, and so on. Six files also
carry planted data problems (duplicate keys, missing keys, out-of-range values, unknown codes); the rows that carry
them are listed in the manifest as the rows a correct pipeline must quarantine.

All businesses and people are invented. Run:  python datasets/build.py
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import random
import shutil
from collections.abc import Callable
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parent
FILES, GOLD, DRIFT, EXTRA = HERE / "files", HERE / "gold", HERE / "drift", HERE / "extra"

FIRST = [
    "Sarah",
    "James",
    "Emma",
    "Tom",
    "Claire",
    "Oliver",
    "Grace",
    "Jack",
    "Sophie",
    "Liam",
    "Hannah",
    "Ben",
    "Chloe",
    "Ryan",
    "Lucy",
    "Daniel",
    "Amelia",
    "Sam",
    "Megan",
    "Luke",
    "Ella",
    "Josh",
    "Kate",
    "Nathan",
]
LAST = [
    "Clarke",
    "Walker",
    "Bennett",
    "Hughes",
    "Turner",
    "Parker",
    "Morgan",
    "Reid",
    "Fletcher",
    "Hayes",
    "Mitchell",
    "Foster",
    "Bailey",
    "Graham",
    "Spencer",
    "Cooper",
    "Wallace",
    "Harper",
    "Lawson",
    "Brooks",
]
COMPANIES = [
    "Acme Home & Garden",
    "Tidewater Landscapes",
    "Kowhai Builders",
    "Bayline Cafe Group",
    "Northgate Property",
    "Pukeko Joinery",
    "Riverbend Motors",
    "Seaview Dental",
    "Ironleaf Studio",
    "Harbourside Events",
]


def money(x: float) -> float:
    return round(x + 0.0, 2)


def write_csv(
    path: Path,
    rows: list[list[Any]],
    delim: str = ",",
    encoding: str = "utf-8",
    pre: list[str] | None = None,
    post: list[str] | None = None,
) -> None:
    buf = io.StringIO()
    for line in pre or []:
        buf.write(line + "\r\n")
    w = csv.writer(buf, delimiter=delim, lineterminator="\r\n")
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    for line in post or []:
        buf.write(line + "\r\n")
    path.write_bytes(buf.getvalue().encode(encoding))


def write_fixed(
    path: Path,
    header: list[str],
    rows: list[list[Any]],
    widths: list[int],
    pre: list[str] | None = None,
    post: list[str] | None = None,
    right: set[int] | None = None,
) -> None:
    right = right or set()

    def fmt(cells: list[Any]) -> str:
        out = []
        for i, (c, w) in enumerate(zip(cells, widths, strict=True)):
            s = "" if c is None else str(c)
            out.append(s.rjust(w) if i in right else s.ljust(w))
        return " ".join(out).rstrip()

    lines = (pre or []) + [fmt(header)] + [fmt(r) for r in rows] + (post or [])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def new_book() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Acme export"
    wb.properties.created = dt.datetime(2026, 4, 1)
    wb.properties.modified = dt.datetime(2026, 4, 1)
    return wb


def merged_header(ws: Any, row: int, groups: list[tuple[str, list[str]]]) -> None:
    """Two header rows: a group label merged across its sub-columns, or one label merged down both rows."""
    col = 1
    for label, subs in groups:
        if subs:
            ws.cell(row, col, label)
            if len(subs) > 1:
                ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + len(subs) - 1)
            for i, s in enumerate(subs):
                ws.cell(row + 1, col + i, s).font = Font(bold=True)
            ws.cell(row, col).alignment = Alignment(horizontal="center")
            col += len(subs)
        else:
            ws.cell(row, col, label)
            ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col)
            col += 1
        ws.cell(row, col - 1).font = Font(bold=True)


# =============================================================== gold tables

COLUMNS = {
    "bank_transactions": ["account", "txn_date", "description", "amount", "balance", "reference"],
    "supplier_prices": ["supplier", "sku", "description", "unit", "pack_size", "price", "currency", "valid_from"],
    "pos_sales": ["order_id", "line_no", "sold_at", "store", "sku", "qty", "unit_price", "discount", "line_total"],
    "timesheets": ["employee_id", "employee_name", "work_date", "project_code", "hours_normal", "hours_overtime"],
    "sensor_readings": ["device_id", "read_at", "temperature_c", "humidity_pct", "battery_v"],
    "crm_contacts": ["contact_id", "full_name", "email", "phone", "company", "created_on", "lifetime_value"],
}
KEYS = {
    "bank_transactions": ["account", "reference"],
    "supplier_prices": ["supplier", "sku"],
    "pos_sales": ["order_id", "line_no"],
    "timesheets": ["employee_id", "work_date", "project_code"],
    "sensor_readings": ["device_id", "read_at"],
    "crm_contacts": ["contact_id"],
}
DECIMALS = {"amount", "balance", "price", "unit_price", "discount", "line_total", "lifetime_value"}


def gold_text(col: str, v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, dt.date):
        return v.isoformat()
    if col in DECIMALS:
        return f"{v:.2f}"
    if isinstance(v, float):
        return repr(round(v, 4))
    return str(v)


def write_gold(path: Path, table: str, records: list[dict[str, Any]]) -> None:
    cols = COLUMNS[table]
    write_csv(path, [cols] + [[gold_text(c, r[c]) for c in cols] for r in records])


# =============================================================== sources

D0 = dt.date(2026, 3, 1)


def bank_records(
    rng: random.Random, account: str, start: dt.date, n: int = 28, opening: float = 8000.0
) -> list[dict[str, Any]]:
    payees_out = [
        "Acme Fasteners Ltd",
        "Harbour Power",
        "Rent - Lighthouse Properties",
        "Bayline Plumbing Supply",
        "Café Kahu",
        "Fuel - Tidewater Services",
        "Kowhai Electrical",
        "Insurance - Pohutukawa Mutual",
    ]
    payees_in = [
        "Eftpos takings Store 1",
        "Eftpos takings Store 2",
        "Online orders payout",
        "Invoice 2231 Kowhai Builders",
    ]
    out, bal, ref = [], opening, rng.randint(400000, 600000)
    for i in range(n):
        day = start + dt.timedelta(days=i * 27 // n)
        if rng.random() < 0.45:
            amt = money(rng.uniform(300, 4200))
            desc = rng.choice(payees_in)
        else:
            amt = -money(rng.uniform(8, 2600))
            desc = rng.choice(payees_out)
        bal = money(bal + amt)
        ref += rng.randint(1, 40)
        out.append(
            {
                "account": account,
                "txn_date": day,
                "description": desc,
                "amount": amt,
                "balance": bal,
                "reference": f"TXN{ref}",
            }
        )
    return out


def supplier_records(
    rng: random.Random,
    supplier: str,
    currency: str,
    valid: dt.date,
    prefix: str,
    n: int = 24,
    units: tuple[str, ...] = ("EA", "BOX", "PK"),
) -> list[dict[str, Any]]:
    words = ["Galvanised", "Stainless", "Treated", "Heavy duty", "Compact", "Outdoor", "Brass", "Coated"]
    things = [
        "screw 8g x 30mm",
        "bolt M10 x 80",
        "wall anchor",
        "hinge 75mm",
        "bracket",
        "tap washer",
        "cable tie",
        "deck screw",
        "hose clamp",
        "wire rope 4mm",
        "sealant 300ml",
        "pine batten",
    ]
    out, seen = [], set()
    while len(out) < n:
        sku = f"{prefix}-{rng.randint(1000, 9999)}"
        if sku in seen:
            continue
        seen.add(sku)
        unit = rng.choice(units)
        out.append(
            {
                "supplier": supplier,
                "sku": sku,
                "description": f"{rng.choice(words)} {rng.choice(things)}",
                "unit": unit,
                "pack_size": 1 if unit == "EA" else rng.choice([10, 20, 50, 100, 200]),
                "price": money(rng.uniform(0.4, 380)),
                "currency": currency,
                "valid_from": valid,
            }
        )
    return out


def pos_records(
    rng: random.Random, store: str, day: dt.date, n_orders: int = 14, start_no: int = 100200
) -> list[dict[str, Any]]:
    out = []
    t = dt.datetime.combine(day, dt.time(8, 30))
    for k in range(n_orders):
        t += dt.timedelta(minutes=rng.randint(3, 45), seconds=rng.randint(0, 59))
        oid = f"R{start_no + k * 7}"
        for line in range(1, rng.randint(1, 4) + 1):
            qty = -1 if rng.random() < 0.05 else rng.randint(1, 5)
            price = money(rng.choice([4.99, 12.5, 19.9, 2.49, 89.0, 34.95, 7.5, 149.0]))
            disc = money(rng.choice([0, 0, 0, 1.0, 2.5, 5.0])) if qty > 0 else 0.0
            disc = min(disc, money(qty * price)) if qty > 0 else 0.0
            out.append(
                {
                    "order_id": oid,
                    "line_no": line,
                    "sold_at": t,
                    "store": store,
                    "sku": f"HG-{rng.randint(10000, 99999)}",
                    "qty": qty,
                    "unit_price": price,
                    "discount": disc,
                    "line_total": money(qty * price - disc),
                }
            )
    return out


EMPLOYEES = [(f"E{1001 + i}", FIRST[i], LAST[i]) for i in range(8)]
PROJECTS = ["P-1001", "P-1002", "P-2040", "P-3100", "ADMIN"]


def timesheet_records(rng: random.Random, start: dt.date, days: int = 5, people: int = 6) -> list[dict[str, Any]]:
    out = []
    for eid, first, last in EMPLOYEES[:people]:
        for d in range(days):
            day = start + dt.timedelta(days=d)
            normal = rng.choice([7.5, 8.0, 6.0, 7.25, 4.5, 8.0])
            ot = rng.choice([0.0, 0.0, 0.0, 1.5, 2.0, 0.75])
            out.append(
                {
                    "employee_id": eid,
                    "employee_name": f"{first} {last}",
                    "work_date": day,
                    "project_code": rng.choice(PROJECTS),
                    "hours_normal": normal,
                    "hours_overtime": ot,
                }
            )
    return out


def sensor_records(rng: random.Random, start: dt.datetime, devices: list[str], per: int = 12) -> list[dict[str, Any]]:
    out = []
    base = {d: (rng.choice([-18.0, 3.0, 20.0]), rng.uniform(3.5, 3.9)) for d in devices}
    for i in range(per):
        ts = start + dt.timedelta(minutes=15 * i)
        for d in devices:
            t0, b0 = base[d]
            out.append(
                {
                    "device_id": d,
                    "read_at": ts,
                    "temperature_c": round(t0 + rng.uniform(-1.5, 1.5), 1),
                    "humidity_pct": round(rng.uniform(40, 95), 1),
                    "battery_v": round(b0 - 0.002 * i, 3),
                }
            )
    return out


def phone_e164(rng: random.Random) -> str:
    if rng.random() < 0.5:
        return f"+649{rng.randint(2000000, 9999999)}"
    return f"+6421{rng.randint(1000000, 9999999)}"


def crm_records(rng: random.Random, n: int = 26, start_id: int = 1200) -> list[dict[str, Any]]:
    out, used = [], set()
    for i in range(n):
        while True:
            first, last = rng.choice(FIRST), rng.choice(LAST)
            if (first, last) not in used:
                used.add((first, last))
                break
        out.append(
            {
                "contact_id": f"C-{start_id + i * 3:05d}",
                "full_name": f"{first} {last}",
                "first": first,
                "last": last,
                "email": f"{first.lower()}.{last.lower()}@example.com" if rng.random() > 0.1 else None,
                "phone": phone_e164(rng) if rng.random() > 0.1 else None,
                "company": rng.choice(COMPANIES),
                "created_on": dt.date(2025, 1, 1) + dt.timedelta(days=rng.randint(0, 430)),
                "lifetime_value": money(rng.uniform(0, 12000)),
            }
        )
    return out


def nz_local(e164: str | None, style: str) -> str:
    if not e164:
        return ""
    rest = e164[3:]
    if style == "intl":
        return (
            f"+64 {rest[0]} {rest[1:4]} {rest[4:]}"
            if rest.startswith("9")
            else f"+64 {rest[:2]} {rest[2:5]} {rest[5:]}"
        )
    if style == "local":
        return f"(0{rest[0]}) {rest[1:4]} {rest[4:]}" if rest.startswith("9") else f"0{rest[:2]} {rest[2:5]} {rest[5:]}"
    return f"64 {rest[:2]} {rest[2:5]} {rest[5:]}"


def eu_num(x: float, thousands: bool = True) -> str:
    s = f"{x:,.2f}" if thousands else f"{x:.2f}"
    return s.replace(",", "_").replace(".", ",").replace("_", ".")


def usd(x: float, parens: bool = False, symbol: str = "$") -> str:
    s = f"{symbol}{abs(x):,.2f}"
    if x < 0:
        return f"({s})" if parens else f"-{s}"
    return s


# =============================================================== file renderers

Manifest = dict[str, Any]


def bank_files(add: Callable[..., None]) -> None:
    rng = random.Random(11)
    acct1, acct2 = "12-3456-0123456-00", "38-9000-7654321-01"

    recs = bank_records(rng, acct1, D0)
    rows = [["Date", "Reference", "Details", "Amount", "Balance"]]
    rows += [
        [
            r["txn_date"].strftime("%d/%m/%Y"),
            r["reference"],
            r["description"],
            f"{r['amount']:.2f}",
            f"{r['balance']:.2f}",
        ]
        for r in recs
    ]
    write_csv(
        FILES / "harbour_bank_2026-03.csv",
        rows,
        pre=["Harbour Bank - Statement", f"Account: {acct1}", "Period: 01/03/2026 to 31/03/2026", ""],
        post=[f"Closing Balance,,,,{recs[-1]['balance']:.2f}"],
    )
    add(
        "b1",
        "harbour_bank_2026-03.csv",
        "harbour_bank",
        "bank_transactions",
        recs,
        ["preamble lines", "account number only in the preamble", "dd/mm/yyyy dates", "closing-balance footer"],
    )

    recs = bank_records(rng, acct2, D0)
    rows = [["Transaction Date", "Description", "Debit", "Credit", "Balance", "Ref No", "Account Number"]]
    rows += [
        [
            r["txn_date"].strftime("%d %b %Y"),
            r["description"],
            usd(-r["amount"]) if r["amount"] < 0 else "",
            usd(r["amount"]) if r["amount"] > 0 else "",
            usd(r["balance"]),
            r["reference"],
            r["account"],
        ]
        for r in recs
    ]
    write_csv(FILES / "kauri_bank_2026-03.csv", rows)
    add(
        "b2",
        "kauri_bank_2026-03.csv",
        "kauri_bank",
        "bank_transactions",
        recs,
        ["separate debit and credit columns", "$1,234.56 amounts", "'03 Mar 2026' dates"],
    )

    recs = bank_records(rng, acct1, dt.date(2026, 2, 1))
    rows = [["Value Date", "Narrative", "Amount", "Running Balance", "Transaction ID", "Acct"]]
    rows += [
        [
            r["txn_date"].strftime("%d.%m.%Y"),
            r["description"],
            eu_num(r["amount"]),
            eu_num(r["balance"]),
            r["reference"],
            r["account"],
        ]
        for r in recs
    ]
    write_csv(FILES / "harbour_bank_export_eu.csv", rows, delim=";", encoding="cp1252")
    add(
        "b3",
        "harbour_bank_export_eu.csv",
        "harbour_bank_eu",
        "bank_transactions",
        recs,
        ["Windows-1252 encoding", "semicolon delimiter", "decimal commas (1.234,56)", "dd.mm.yyyy dates"],
    )

    recs = bank_records(rng, acct2, dt.date(2026, 1, 1), n=30)
    wb = new_book()
    ws = wb.create_sheet("Transactions")
    ws.append(["Kauri Bank - Transaction export"])
    ws.append([f"Account No: {acct2}"])
    ws.append([])
    ws.append(["Posted Date", "Particulars", "Amount", "Balance", "Txn ID"])
    for r in recs:
        ws.append(
            [dt.datetime.combine(r["txn_date"], dt.time()), r["description"], r["amount"], r["balance"], r["reference"]]
        )
    wb.save(FILES / "kauri_bank_q1.xlsx")
    add(
        "b4",
        "kauri_bank_q1.xlsx",
        "kauri_bank_xlsx",
        "bank_transactions",
        recs,
        ["title rows above the header", "account number in a title row", "real Excel dates"],
    )

    recs = bank_records(rng, "02-0100-0999888-00", dt.date(2026, 3, 1))
    rows = [["Date", "Memo", "Amount", "Balance", "Reference", "Account"]]
    rows += [
        [
            r["txn_date"].strftime("%m/%d/%Y"),
            r["description"],
            usd(r["amount"], parens=True, symbol=""),
            usd(r["balance"], parens=True, symbol=""),
            r["reference"],
            r["account"],
        ]
        for r in recs
    ]
    write_csv(FILES / "harbour_bank_usd.tsv", rows, delim="\t")
    add(
        "b5",
        "harbour_bank_usd.tsv",
        "harbour_bank_us",
        "bank_transactions",
        recs,
        ["tab delimiter", "mm/dd/yyyy dates", "(45.20) negatives", "thousands separators"],
    )

    recs = bank_records(rng, acct1, dt.date(2026, 4, 1))
    bad = [dict(recs[5]), dict(recs[11], reference=None), dict(recs[17], txn_date=dt.date(2031, 4, 18))]
    file_recs = recs[:6] + [bad[0]] + recs[6:11] + [bad[1]] + recs[12:17] + [bad[2]] + recs[18:]
    rows = [["Date", "Reference", "Details", "Amount", "Balance"]]
    rows += [
        [
            r["txn_date"].strftime("%d/%m/%Y"),
            r["reference"] or "",
            r["description"],
            f"{r['amount']:.2f}",
            f"{r['balance']:.2f}",
        ]
        for r in file_recs
    ]
    write_csv(
        FILES / "harbour_bank_2026-04.csv",
        rows,
        pre=["Harbour Bank - Statement", f"Account: {acct1}", "Period: 01/04/2026 to 30/04/2026", ""],
        post=[f"Closing Balance,,,,{recs[-1]['balance']:.2f}"],
    )
    gold = [r for i, r in enumerate(recs) if i not in (11, 17)]
    add(
        "b6",
        "harbour_bank_2026-04.csv",
        "harbour_bank",
        "bank_transactions",
        gold,
        ["preamble lines", "closing-balance footer"],
        planted=[
            {"problem": "duplicate reference", "test": "unique", "rows": 1},
            {"problem": "missing reference", "test": "not_null", "rows": 1},
            {"problem": "date in 2031", "test": "range", "rows": 1},
        ],
        quarantine=3,
    )


def supplier_files(add: Callable[..., None]) -> None:
    rng = random.Random(22)
    valid = dt.date(2026, 4, 1)

    recs = supplier_records(rng, "Acme Fasteners Ltd", "NZD", valid, "AF", n=27)
    wb = new_book()
    notes = wb.create_sheet("Read me")
    notes.append(["Acme Fasteners Ltd - trade price list"])
    notes.append(["Prices exclude GST. Call 0800 555 010 for volume pricing."])
    sheets = {"Screws": recs[:10], "Bolts": recs[10:19], "Anchors": recs[19:]}
    for name, part in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(["Supplier: Acme Fasteners Ltd"])
        ws.append(["Effective from: 2026-04-01"])
        ws.append([])
        merged_header(
            ws,
            4,
            [
                ("Product", ["Code", "Description"]),
                ("Pack", ["Unit", "Qty"]),
                ("Pricing", ["Price ex GST", "Currency"]),
            ],
        )
        for r in part:
            ws.append([r["sku"], r["description"], r["unit"], r["pack_size"], r["price"], r["currency"]])
    wb.save(FILES / "acme_fasteners_2026.xlsx")
    add(
        "s1",
        "acme_fasteners_2026.xlsx",
        "acme_fasteners",
        "supplier_prices",
        recs,
        ["three data sheets plus a notes sheet", "merged two-row header", "supplier and date only in title rows"],
    )

    recs = supplier_records(rng, "Bayline Plumbing Supply", "NZD", valid, "BP", units=("EA", "BOX", "M"))
    rows = [["Vendor", "Item Code", "Item Name", "UOM", "Pack", "Unit Price", "Effective"]]
    rows += [
        [
            r["supplier"],
            r["sku"],
            r["description"],
            r["unit"].lower(),
            r["pack_size"],
            f"NZ${r['price']:,.2f}",
            r["valid_from"].isoformat(),
        ]
        for r in recs
    ]
    write_csv(FILES / "bayline_prices.csv", rows)
    add(
        "s2",
        "bayline_prices.csv",
        "bayline",
        "supplier_prices",
        recs,
        ["NZ$ prices", "lower-case units", "currency only implied by the symbol"],
    )

    recs = supplier_records(rng, "Nordwerk Tools GmbH", "EUR", valid, "NW", units=("EA", "PK", "BOX"))
    rows = [["Artikel", "Bezeichnung", "Einheit", "VPE", "Preis", "Gueltig ab"]]
    rows += [
        [
            r["sku"],
            r["description"],
            r["unit"],
            r["pack_size"],
            f"{eu_num(r['price'])} €",
            r["valid_from"].strftime("%d.%m.%Y"),
        ]
        for r in recs
    ]
    write_csv(FILES / "nordwerk_preisliste.csv", rows, delim=";", pre=["Lieferant: Nordwerk Tools GmbH"])
    add(
        "s3",
        "nordwerk_preisliste.csv",
        "nordwerk",
        "supplier_prices",
        recs,
        [
            "German headers",
            "semicolon delimiter",
            "decimal commas with a euro sign",
            "dd.mm.yyyy dates",
            "supplier name only in the first line",
        ],
    )

    recs = supplier_records(rng, "Kowhai Electrical", "NZD", valid, "KE", units=("EA", "M", "BOX"))
    write_fixed(
        FILES / "kowhai_erp_export.txt",
        ["SUPPLIER", "ITEM_NO", "DESCRIPTION", "UOM", "PK", "PRICE", "CUR", "EFF_DATE"],
        [
            [
                r["supplier"],
                r["sku"],
                r["description"],
                r["unit"],
                r["pack_size"],
                f"{r['price']:.2f}",
                r["currency"],
                r["valid_from"].strftime("%Y%m%d"),
            ]
            for r in recs
        ],
        [18, 8, 28, 4, 4, 9, 4, 8],
        right={4, 5},
    )
    add(
        "s4",
        "kowhai_erp_export.txt",
        "kowhai",
        "supplier_prices",
        recs,
        ["fixed-width layout", "descriptions with spaces", "yyyymmdd dates"],
    )

    recs = supplier_records(rng, "Ironleaf Tools Pty", "AUD", valid, "IL")
    doc = {
        "supplier": {"name": "Ironleaf Tools Pty", "country": "AU"},
        "valid_from": "2026-04-01",
        "products": [
            {
                "code": r["sku"],
                "title": r["description"],
                "pack": {"unit": r["unit"], "size": r["pack_size"]},
                "price": {"amount": r["price"], "currency": r["currency"]},
            }
            for r in recs
        ],
    }
    (FILES / "ironleaf_catalogue.json").write_text(json.dumps(doc, indent=2), encoding="utf-8")
    add(
        "s5",
        "ironleaf_catalogue.json",
        "ironleaf",
        "supplier_prices",
        recs,
        ["nested JSON", "supplier and date at the top level"],
    )

    recs = supplier_records(rng, "Tidewater Timber", "NZD", valid, "TT", units=("M", "EA", "PK"))
    file_recs = [dict(r) for r in recs]
    file_recs.insert(8, dict(recs[7], price=money(recs[7]["price"] + 1)))
    file_recs[14]["currency"] = "XYZ"
    file_recs[20]["price"] = -5.0
    rows = [["Supplier", "SKU", "Description", "Unit", "Pack Size", "Price", "Currency", "Valid From"]]
    rows += [
        [
            r["supplier"],
            r["sku"],
            r["description"],
            r["unit"],
            r["pack_size"],
            f"{r['price']:.2f}",
            r["currency"],
            r["valid_from"].isoformat(),
        ]
        for r in file_recs
    ]
    write_csv(FILES / "tidewater_timber.csv", rows)
    gold = [r for i, r in enumerate(recs) if i not in (13, 19)]
    add(
        "s6",
        "tidewater_timber.csv",
        "tidewater",
        "supplier_prices",
        gold,
        ["clean layout"],
        planted=[
            {"problem": "duplicate SKU", "test": "unique", "rows": 1},
            {"problem": "unknown currency XYZ", "test": "accepted_values", "rows": 1},
            {"problem": "negative price", "test": "range", "rows": 1},
        ],
        quarantine=3,
    )


def pos_files(add: Callable[..., None]) -> None:
    rng = random.Random(33)

    def orders(recs: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
        out: dict[str, list[dict[str, Any]]] = {}
        for r in recs:
            out.setdefault(r["order_id"], []).append(r)
        return out

    recs = pos_records(rng, "AK01", dt.date(2026, 3, 2))
    doc = {
        "export": {
            "generated_at": "2026-03-02T23:00:05",
            "store": "AK01",
            "orders": [
                {
                    "order_no": oid,
                    "created": ls[0]["sold_at"].isoformat(),
                    "cashier": rng.choice(FIRST),
                    "items": [
                        {
                            "line": x["line_no"],
                            "sku": x["sku"],
                            "qty": x["qty"],
                            "price": x["unit_price"],
                            "discount": x["discount"],
                            "total": x["line_total"],
                        }
                        for x in ls
                    ],
                }
                for oid, ls in orders(recs).items()
            ],
        }
    }
    (FILES / "pos_export_2026-03-02.json").write_text(json.dumps(doc, indent=2), encoding="utf-8")
    add(
        "p1",
        "pos_export_2026-03-02.json",
        "pos_json",
        "pos_sales",
        recs,
        ["nested JSON orders with line items", "store code only at the top level", "a cashier field to ignore"],
    )

    recs = pos_records(rng, "AK02", dt.date(2026, 3, 3), start_no=200100)
    lines = []
    for oid, ls in orders(recs).items():
        items = []
        for x in ls:
            it: dict[str, Any] = {"sku": x["sku"], "qty": x["qty"], "unit": x["unit_price"], "total": x["line_total"]}
            if x["discount"]:
                it["disc"] = x["discount"]
            items.append(it)
        lines.append(
            json.dumps(
                {
                    "receipt": oid,
                    "ts": ls[0]["sold_at"].strftime("%Y-%m-%d %H:%M:%S"),
                    "store": {"code": "ak02"},
                    "lines": items,
                }
            )
        )
    (FILES / "pos_ak02_2026-03-03.ndjson").write_text("\n".join(lines) + "\n", encoding="utf-8")
    add(
        "p2",
        "pos_ak02_2026-03-03.ndjson",
        "pos_ndjson",
        "pos_sales",
        recs,
        [
            "one JSON object per line",
            "line number only implied by position",
            "discount missing when zero",
            "lower-case store code",
        ],
    )

    recs = pos_records(rng, "NS01", dt.date(2026, 3, 4), start_no=300050)
    rows = [["Receipt No", "Line", "Sold At", "Store", "PLU", "Quantity", "Each", "Disc", "Line Amount"]]
    rows += [
        [
            r["order_id"],
            r["line_no"],
            r["sold_at"].strftime("%d/%m/%Y %H:%M:%S"),
            r["store"],
            r["sku"],
            r["qty"],
            usd(r["unit_price"]),
            usd(r["discount"]),
            usd(r["line_total"]),
        ]
        for r in recs
    ]
    write_csv(FILES / "pos_flat_export.csv", rows)
    add("p3", "pos_flat_export.csv", "pos_flat", "pos_sales", recs, ["$ amounts", "dd/mm/yyyy hh:mm:ss times"])

    recs1 = pos_records(rng, "AK01", dt.date(2026, 3, 5), n_orders=8, start_no=400000)
    recs2 = pos_records(rng, "AK02", dt.date(2026, 3, 5), n_orders=8, start_no=410000)
    wb = new_book()
    for name, part in (("AK01", recs1), ("AK02", recs2)):
        ws = wb.create_sheet(name)
        ws.append(["Order", "Line #", "Date", "Time", "SKU", "Qty", "Unit Price", "Discount", "Total"])
        for r in part:
            ws.append(
                [
                    r["order_id"],
                    r["line_no"],
                    dt.datetime.combine(r["sold_at"].date(), dt.time()),
                    r["sold_at"].strftime("%H:%M:%S"),
                    r["sku"],
                    r["qty"],
                    r["unit_price"],
                    r["discount"],
                    r["line_total"],
                ]
            )
    summary = wb.create_sheet("Summary")
    summary.append(["Store", "Sales"])
    summary.append(["AK01", round(sum(r["line_total"] for r in recs1), 2)])
    summary.append(["AK02", round(sum(r["line_total"] for r in recs2), 2)])
    wb.save(FILES / "pos_weekly.xlsx")
    add(
        "p4",
        "pos_weekly.xlsx",
        "pos_xlsx",
        "pos_sales",
        recs1 + recs2,
        [
            "one sheet per store",
            "store only in the sheet name",
            "date and time in separate columns",
            "a summary sheet to skip",
        ],
    )

    recs = pos_records(rng, "MK01", dt.date(2026, 3, 6), start_no=500300)
    doc2 = [
        {
            "id": oid,
            "store_id": "MK01",
            "ts": ls[0]["sold_at"].strftime("%Y%m%dT%H%M%S"),
            "lines": [
                {
                    "n": x["line_no"],
                    "sku": x["sku"],
                    "qty": x["qty"],
                    "price_cents": round(x["unit_price"] * 100),
                    "discount_cents": round(x["discount"] * 100),
                    "total_cents": round(x["line_total"] * 100),
                }
                for x in ls
            ],
        }
        for oid, ls in orders(recs).items()
    ]
    (FILES / "pos_cents.json").write_text(json.dumps(doc2), encoding="utf-8")
    add(
        "p5",
        "pos_cents.json",
        "pos_cents",
        "pos_sales",
        recs,
        ["amounts in cents", "compact 20260306T091522 timestamps", "minified JSON"],
    )

    recs = pos_records(rng, "AK01", dt.date(2026, 3, 9), start_no=600000)
    by = orders(recs)
    ids = list(by)
    bad_null, bad_qty = ids[3], ids[7]
    items_out = []
    for oid, ls in by.items():
        items = [
            {
                "line": x["line_no"],
                "sku": x["sku"],
                "qty": 999 if (oid == bad_qty and x["line_no"] == 1) else x["qty"],
                "price": x["unit_price"],
                "discount": x["discount"],
                "total": x["line_total"],
            }
            for x in ls
        ]
        if oid == ids[10]:
            items.append(dict(items[0]))
        items_out.append(
            {
                "order_no": None if oid == bad_null else oid,
                "created": ls[0]["sold_at"].isoformat(),
                "cashier": rng.choice(FIRST),
                "items": items,
            }
        )
    doc = {"export": {"generated_at": "2026-03-09T23:00:05", "store": "AK01", "orders": items_out}}
    (FILES / "pos_export_2026-03-09.json").write_text(json.dumps(doc, indent=2), encoding="utf-8")
    gold = [r for r in recs if r["order_id"] != bad_null and not (r["order_id"] == bad_qty and r["line_no"] == 1)]
    add(
        "p6",
        "pos_export_2026-03-09.json",
        "pos_json",
        "pos_sales",
        gold,
        ["nested JSON orders"],
        planted=[
            {"problem": "order with no order number", "test": "not_null", "rows": len(by[bad_null])},
            {"problem": "quantity 999", "test": "range", "rows": 1},
            {"problem": "line item sent twice", "test": "unique", "rows": 1},
        ],
        quarantine=len(by[bad_null]) + 2,
    )


def timesheet_files(add: Callable[..., None]) -> None:
    rng = random.Random(44)
    names = {e: (f, la) for e, f, la in EMPLOYEES}

    recs = timesheet_records(rng, dt.date(2026, 3, 16))
    wb = new_book()
    ws = wb.create_sheet("Week 12")
    ws.append(["Acme Home & Garden - timesheet week 12"])
    ws.append([])
    merged_header(
        ws, 3, [("Employee", ["ID", "Name"]), ("Date", []), ("Project", []), ("Hours", ["Normal", "Overtime"])]
    )
    for r in recs:
        f, la = names[r["employee_id"]]
        ws.append(
            [
                r["employee_id"],
                f"{la}, {f}",
                dt.datetime.combine(r["work_date"], dt.time()),
                r["project_code"],
                r["hours_normal"],
                r["hours_overtime"],
            ]
        )
    wb.save(FILES / "timesheet_week12.xlsx")
    add(
        "t1",
        "timesheet_week12.xlsx",
        "timesheet_xlsx",
        "timesheets",
        recs,
        ["merged two-row header", "names written 'Last, First'", "a title row"],
    )

    def hm(x: float) -> str:
        return f"{int(x)}:{round((x - int(x)) * 60):02d}"

    recs = timesheet_records(rng, dt.date(2026, 3, 23))
    rows = [["Emp #", "Staff Name", "Shift Date", "Job Code", "Ordinary", "OT"]]
    rows += [
        [
            r["employee_id"],
            r["employee_name"],
            r["work_date"].isoformat(),
            r["project_code"].lower(),
            hm(r["hours_normal"]),
            hm(r["hours_overtime"]),
        ]
        for r in recs
    ]
    write_csv(FILES / "timesheets_march.csv", rows)
    add(
        "t2",
        "timesheets_march.csv",
        "timesheet_hm",
        "timesheets",
        recs,
        ["hours written as h:mm", "lower-case project codes"],
    )

    recs = timesheet_records(rng, dt.date(2026, 3, 9))
    write_fixed(
        FILES / "payroll_export.txt",
        ["EMPNO", "NAME", "WORKDATE", "PROJ", "NRM", "OVT"],
        [
            [
                r["employee_id"],
                r["employee_name"],
                r["work_date"].strftime("%Y%m%d"),
                r["project_code"],
                f"{r['hours_normal']:.2f}",
                f"{r['hours_overtime']:.2f}",
            ]
            for r in recs
        ],
        [6, 18, 8, 7, 5, 5],
        pre=["Exported: 2026-03-14"],
        right={4, 5},
        post=[f"END OF FILE {len(recs)} RECORDS"],
    )
    add(
        "t3",
        "payroll_export.txt",
        "payroll_fixed",
        "timesheets",
        recs,
        ["fixed-width layout", "yyyymmdd dates", "an export-date line and an end-of-file line"],
    )

    recs = timesheet_records(rng, dt.date(2026, 3, 2))
    rows = [["Employee ID", "Employee", "Date", "Project", "Hours", "Overtime Hours"]]
    rows += [
        [
            r["employee_id"],
            r["employee_name"],
            r["work_date"].strftime("%d/%m/%Y"),
            r["project_code"],
            eu_num(r["hours_normal"], False),
            eu_num(r["hours_overtime"], False),
        ]
        for r in recs
    ]
    write_csv(FILES / "timesheet_bom.csv", rows, delim=";", encoding="utf-8-sig")
    add(
        "t4",
        "timesheet_bom.csv",
        "timesheet_bom",
        "timesheets",
        recs,
        ["UTF-8 byte-order mark", "semicolon delimiter", "decimal commas"],
    )

    recs = timesheet_records(rng, dt.date(2026, 2, 23))
    emp: dict[str, dict[str, Any]] = {}
    for r in recs:
        f, la = names[r["employee_id"]]
        e = emp.setdefault(r["employee_id"], {"id": r["employee_id"], "name": {"first": f, "last": la}, "entries": []})
        e["entries"].append(
            {
                "date": r["work_date"].isoformat(),
                "project": r["project_code"],
                "normal": r["hours_normal"],
                "overtime": r["hours_overtime"],
            }
        )
    (FILES / "timesheets.json").write_text(
        json.dumps({"period": "2026-W09", "employees": list(emp.values())}, indent=1), encoding="utf-8"
    )
    add(
        "t5",
        "timesheets.json",
        "timesheet_json",
        "timesheets",
        recs,
        ["nested JSON", "first and last name in separate fields"],
    )

    recs = timesheet_records(rng, dt.date(2026, 3, 30))
    file_recs = [dict(r) for r in recs]
    file_recs[4]["hours_normal"] = 25.0
    file_recs.insert(10, dict(recs[9]))
    file_recs[18]["employee_id"] = None
    rows = [["Employee ID", "Employee Name", "Work Date", "Project Code", "Normal Hours", "Overtime Hours"]]
    rows += [
        [
            r["employee_id"] or "",
            r["employee_name"],
            r["work_date"].isoformat(),
            r["project_code"],
            r["hours_normal"],
            r["hours_overtime"],
        ]
        for r in file_recs
    ]
    write_csv(FILES / "timesheet_week13.csv", rows)
    gold = [r for i, r in enumerate(recs) if i not in (4, 17)]
    add(
        "t6",
        "timesheet_week13.csv",
        "timesheet_clean",
        "timesheets",
        gold,
        ["clean layout"],
        planted=[
            {"problem": "25 hours in a day", "test": "range", "rows": 1},
            {"problem": "entry sent twice", "test": "unique", "rows": 1},
            {"problem": "missing employee ID", "test": "not_null", "rows": 1},
        ],
        quarantine=3,
    )


def sensor_files(add: Callable[..., None]) -> None:
    rng = random.Random(55)
    t0 = dt.datetime(2026, 3, 1, 8, 0)
    epoch = dt.datetime(1970, 1, 1)

    recs = sensor_records(rng, t0, ["CS-01", "CS-02", "CS-03", "CS-04"])
    rows = [["device", "ts", "temp", "rh", "bat"]]
    rows += [
        [
            r["device_id"].lower(),
            int((r["read_at"] - epoch).total_seconds()),
            r["temperature_c"],
            r["humidity_pct"],
            r["battery_v"],
        ]
        for r in recs
    ]
    write_csv(FILES / "coldstore_dump_0301.csv", rows)
    add(
        "i1",
        "coldstore_dump_0301.csv",
        "coldstore_csv",
        "sensor_readings",
        recs,
        ["Unix epoch timestamps", "lower-case device IDs", "terse headers"],
    )

    recs = sensor_records(rng, t0 + dt.timedelta(days=1), ["CS-01", "CS-02", "CS-05"])
    rows = [["Sensor ID", "Recorded At", "Temp (F)", "Humidity (%)", "Battery (V)"]]
    rows += [
        [
            r["device_id"],
            r["read_at"].strftime("%Y-%m-%dT%H:%M:%SZ"),
            round(r["temperature_c"] * 9 / 5 + 32, 1),
            r["humidity_pct"],
            r["battery_v"],
        ]
        for r in recs
    ]
    write_csv(FILES / "coldstore_fahrenheit.csv", rows)
    add(
        "i2",
        "coldstore_fahrenheit.csv",
        "coldstore_f",
        "sensor_readings",
        recs,
        ["temperature in Fahrenheit", "ISO timestamps with Z"],
    )

    recs = sensor_records(rng, t0 + dt.timedelta(days=2), ["GW-11", "GW-12", "GW-13"])
    write_fixed(
        FILES / "gateway_dump.txt",
        ["SERIAL", "TIMESTAMP", "TEMP_C", "RH", "VBAT"],
        [
            [
                r["device_id"],
                r["read_at"].strftime("%Y-%m-%d %H:%M:%S"),
                f"{r['temperature_c']:.1f}",
                f"{r['humidity_pct']:.1f}",
                f"{r['battery_v']:.3f}",
            ]
            for r in recs
        ],
        [8, 19, 6, 5, 5],
        right={2, 3, 4},
    )
    add(
        "i3",
        "gateway_dump.txt",
        "gateway_fixed",
        "sensor_readings",
        recs,
        ["fixed-width layout", "timestamps containing a space"],
    )

    recs = sensor_records(rng, t0 + dt.timedelta(days=3), ["CS-03", "CS-04", "CS-06"])
    lines = [
        json.dumps(
            {
                "dev": r["device_id"],
                "ts": r["read_at"].strftime("%Y-%m-%dT%H:%M:%SZ"),
                "payload": {"t": r["temperature_c"], "h": r["humidity_pct"], "bat_mv": round(r["battery_v"] * 1000)},
            }
        )
        for r in recs
    ]
    (FILES / "sensors.ndjson").write_text("\n".join(lines) + "\n", encoding="utf-8")
    add(
        "i4",
        "sensors.ndjson",
        "sensor_ndjson",
        "sensor_readings",
        recs,
        ["one JSON object per line", "nested payload", "battery in millivolts"],
    )

    recs = sensor_records(rng, t0 + dt.timedelta(days=4), ["CS-01", "CS-07"], per=16)
    rows = [["Device", "Time", "Temperature", "Humidity", "Battery"]]
    rows += [
        [
            r["device_id"],
            r["read_at"].strftime("%d/%m/%Y %H:%M"),
            f"{r['temperature_c']}°C",
            f"{r['humidity_pct']}%",
            f"{r['battery_v']}V",
        ]
        for r in recs
    ]
    write_csv(FILES / "sensor_units.psv", rows, delim="|")
    add(
        "i5",
        "sensor_units.psv",
        "sensor_units",
        "sensor_readings",
        recs,
        ["pipe delimiter", "units written inside the values (21.4°C, 55%, 3.71V)", "dd/mm/yyyy hh:mm"],
    )

    recs = sensor_records(rng, t0 + dt.timedelta(days=7), ["CS-01", "CS-02", "CS-03", "CS-04"])
    file_recs = [dict(r) for r in recs]
    file_recs[6]["temperature_c"] = -999.0
    file_recs[13]["humidity_pct"] = 140.0
    file_recs.insert(21, dict(recs[20]))
    rows = [["device", "ts", "temp", "rh", "bat"]]
    rows += [
        [
            r["device_id"].lower(),
            int((r["read_at"] - epoch).total_seconds()),
            r["temperature_c"],
            r["humidity_pct"],
            r["battery_v"],
        ]
        for r in file_recs
    ]
    write_csv(FILES / "coldstore_dump_0308.csv", rows)
    gold = [r for i, r in enumerate(recs) if i not in (6, 13)]
    add(
        "i6",
        "coldstore_dump_0308.csv",
        "coldstore_csv",
        "sensor_readings",
        gold,
        ["Unix epoch timestamps"],
        planted=[
            {"problem": "-999 error code as a temperature", "test": "range", "rows": 1},
            {"problem": "140% humidity", "test": "range", "rows": 1},
            {"problem": "reading sent twice", "test": "unique", "rows": 1},
        ],
        quarantine=3,
    )


def crm_files(add: Callable[..., None]) -> None:
    rng = random.Random(66)

    recs = crm_records(rng, start_id=1200)
    rows = [["First Name", "Last Name", "Email", "Phone", "Company", "Created", "Lifetime Value", "Contact ID"]]
    rows += [
        [
            r["first"],
            r["last"],
            (r["email"] or "").replace(r["first"].lower(), r["first"]),
            nz_local(r["phone"], "intl"),
            r["company"],
            r["created_on"].isoformat(),
            f"{r['lifetime_value']:.2f}",
            r["contact_id"],
        ]
        for r in recs
    ]
    write_csv(FILES / "crm_export_utf16.txt", rows, delim="\t", encoding="utf-16")
    add(
        "c1",
        "crm_export_utf16.txt",
        "crm_utf16",
        "crm_contacts",
        recs,
        [
            "UTF-16 text export",
            "tab delimiter",
            "first and last name in separate columns",
            "mixed-case emails",
            "+64 9 555 1234 phones",
        ],
    )

    recs = crm_records(rng, start_id=2400)
    notes = [
        "Prefers email.",
        "Met at the trade show,\nasked about bulk pricing.",
        "",
        'VIP, "gold" tier.',
        "Call after 3pm",
    ]
    rows = [["Record ID", "Name", "E-mail", "Mobile", "Organisation", "Date Created", "Total Spend", "Notes"]]
    rows += [
        [
            r["contact_id"],
            r["full_name"],
            r["email"] or "",
            nz_local(r["phone"], "local"),
            r["company"],
            r["created_on"].strftime("%d-%b-%Y"),
            usd(r["lifetime_value"]),
            rng.choice(notes),
        ]
        for r in recs
    ]
    write_csv(FILES / "crm_contacts_notes.csv", rows)
    add(
        "c2",
        "crm_contacts_notes.csv",
        "crm_notes",
        "crm_contacts",
        recs,
        ["a notes column with line breaks and quotes", "(09) 555 1234 phones", "'14-Feb-2026' dates", "$ amounts"],
    )

    recs = crm_records(rng, start_id=3600)
    wb = new_book()
    ws = wb.create_sheet("Contacts")
    ws.append(["Customer ID", "Customer Name", "Email Address", "Telephone", "Company", "Signup Date", "LTV"])
    for r in recs:
        ws.append(
            [
                r["contact_id"],
                r["full_name"],
                r["email"],
                r["phone"],
                r["company"],
                dt.datetime.combine(r["created_on"], dt.time()),
                r["lifetime_value"],
            ]
        )
    n = wb.create_sheet("Notes")
    n.append(["Exported by Hannah Graham"])
    n.append(["Do not share outside the sales team"])
    wb.save(FILES / "crm_contacts.xlsx")
    add("c3", "crm_contacts.xlsx", "crm_xlsx", "crm_contacts", recs, ["a notes sheet to skip", "Excel dates"])

    recs = crm_records(rng, start_id=4800)
    doc = {
        "data": {
            "contacts": [
                {
                    "id": r["contact_id"],
                    "name": {"first": r["first"], "last": r["last"]},
                    "contact": {"email": r["email"], "phone": r["phone"]},
                    "org": {"name": r["company"]},
                    "created_at": f"{r['created_on'].isoformat()}T{rng.randint(8, 18):02d}:"
                    f"{rng.randint(0, 59):02d}:00Z",
                    "stats": {"ltv": r["lifetime_value"], "orders": rng.randint(0, 40)},
                }
                for r in recs
            ]
        },
        "meta": {"page": 1, "pages": 1},
    }
    (FILES / "crm_api_dump.json").write_text(json.dumps(doc, indent=2), encoding="utf-8")
    add(
        "c4",
        "crm_api_dump.json",
        "crm_api",
        "crm_contacts",
        recs,
        [
            "nested JSON",
            "first and last name in separate fields",
            "created as a timestamp, not a date",
            "extra fields to ignore",
        ],
    )

    recs = crm_records(rng, start_id=6000)
    rows = [["ID", "Contact Name", "Mail", "Phone Number", "Employer", "Created Date", "Revenue"]]
    rows += [
        [
            r["contact_id"],
            r["full_name"],
            r["email"] or "",
            nz_local(r["phone"], "bare"),
            r["company"],
            f"{r['created_on'].isoformat()}T10:22:00",
            eu_num(r["lifetime_value"]),
        ]
        for r in recs
    ]
    write_csv(FILES / "crm_semicolon.csv", rows, delim=";")
    add(
        "c5",
        "crm_semicolon.csv",
        "crm_semicolon",
        "crm_contacts",
        recs,
        ["semicolon delimiter", "decimal commas", "64 21 555 1234 phones", "timestamps for a date column"],
    )

    recs = crm_records(rng, start_id=7200)
    file_recs = [dict(r) for r in recs]
    file_recs.insert(6, dict(recs[5], lifetime_value=money(recs[5]["lifetime_value"] + 10)))
    file_recs[12]["contact_id"] = None
    file_recs[20]["created_on"] = dt.date(1899, 12, 31)
    rows = [["Contact ID", "Full Name", "Email", "Phone", "Company", "Created On", "Lifetime Value"]]
    rows += [
        [
            r["contact_id"] or "",
            r["full_name"],
            r["email"] or "",
            r["phone"] or "",
            r["company"],
            r["created_on"].isoformat(),
            f"{r['lifetime_value']:.2f}",
        ]
        for r in file_recs
    ]
    write_csv(FILES / "crm_export_2026-04.csv", rows)
    gold = [r for i, r in enumerate(recs) if i not in (11, 19)]
    add(
        "c6",
        "crm_export_2026-04.csv",
        "crm_clean",
        "crm_contacts",
        gold,
        ["clean layout"],
        planted=[
            {"problem": "contact ID used twice", "test": "unique", "rows": 1},
            {"problem": "missing contact ID", "test": "not_null", "rows": 1},
            {"problem": "created on 1899-12-31", "test": "range", "rows": 1},
        ],
        quarantine=3,
    )


# =============================================================== drift scenarios


def drift_files(addd: Callable[..., None]) -> None:
    rng = random.Random(77)
    acct1, acct2 = "12-3456-0123456-00", "38-9000-7654321-01"

    recs = bank_records(rng, acct1, dt.date(2026, 5, 1))
    rows = [["Date", "Reference", "Details", "Amount (NZD)", "Balance"]]
    rows += [
        [
            r["txn_date"].strftime("%d/%m/%Y"),
            r["reference"],
            r["description"],
            f"{r['amount']:.2f}",
            f"{r['balance']:.2f}",
        ]
        for r in recs
    ]
    write_csv(
        DRIFT / "d01_harbour_bank_2026-05.csv",
        rows,
        pre=["Harbour Bank - Statement", f"Account: {acct1}", "Period: 01/05/2026 to 31/05/2026", ""],
        post=[f"Closing Balance,,,,{recs[-1]['balance']:.2f}"],
    )
    addd("d01", "b1", "d01_harbour_bank_2026-05.csv", recs, ["renamed"], "loaded", "Amount became 'Amount (NZD)'")

    recs = bank_records(rng, acct2, dt.date(2026, 4, 1))
    rows = [["Transaction Date", "Description", "Debit", "Credit", "Balance", "Ref No", "Account Number"]]
    rows += [
        [
            r["txn_date"].isoformat(),
            r["description"],
            usd(-r["amount"]) if r["amount"] < 0 else "",
            usd(r["amount"]) if r["amount"] > 0 else "",
            usd(r["balance"]),
            r["reference"],
            r["account"],
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d02_kauri_bank_2026-04.csv", rows)
    addd(
        "d02",
        "b2",
        "d02_kauri_bank_2026-04.csv",
        recs,
        ["format"],
        "loaded",
        "dates changed from '03 Mar 2026' to 2026-03-03",
    )

    recs = supplier_records(rng, "Bayline Plumbing Supply", "NZD", dt.date(2026, 5, 1), "BP", units=("EA", "BOX", "M"))
    rows = [["Vendor", "Item Code", "Item Name", "UOM", "Pack", "Unit Price", "Discount %", "Effective"]]
    rows += [
        [
            r["supplier"],
            r["sku"],
            r["description"],
            r["unit"].lower(),
            r["pack_size"],
            f"NZ${r['price']:,.2f}",
            rng.choice(["0", "5", "10"]),
            r["valid_from"].isoformat(),
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d03_bayline_prices_may.csv", rows)
    addd(
        "d03",
        "s2",
        "d03_bayline_prices_may.csv",
        recs,
        ["added"],
        "review",
        "a new 'Discount %' column the table has no place for",
    )

    recs = pos_records(rng, "NS01", dt.date(2026, 4, 1), start_no=310000)
    rows = [["Receipt No", "Line", "Sold At", "Store", "PLU", "Each", "Qty", "Disc", "Line Amount"]]
    rows += [
        [
            r["order_id"],
            r["line_no"],
            r["sold_at"].strftime("%d/%m/%Y %H:%M:%S"),
            r["store"],
            r["sku"],
            usd(r["unit_price"]),
            r["qty"],
            usd(r["discount"]),
            usd(r["line_total"]),
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d04_pos_flat_export_apr.csv", rows)
    addd(
        "d04",
        "p3",
        "d04_pos_flat_export_apr.csv",
        recs,
        ["renamed", "reordered"],
        "loaded",
        "'Quantity' became 'Qty' and moved after 'Each'",
    )

    recs = timesheet_records(rng, dt.date(2026, 4, 6))
    rows = [["Emp #", "Staff Name", "Shift Date", "Job Code", "Ordinary", "OT"]]
    rows += [
        [
            r["employee_id"],
            r["employee_name"],
            r["work_date"].isoformat(),
            r["project_code"].lower(),
            f"{r['hours_normal']:.2f}",
            f"{r['hours_overtime']:.2f}",
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d05_timesheets_april.csv", rows)
    addd("d05", "t2", "d05_timesheets_april.csv", recs, ["format"], "loaded", "hours changed from 7:30 to 7.50")

    recs = sensor_records(rng, dt.datetime(2026, 4, 2, 8, 0), ["CS-01", "CS-02", "CS-05"])
    rows = [["Sensor ID", "Recorded At", "Temp (C)", "Humidity (%)", "Battery (V)"]]
    rows += [
        [
            r["device_id"],
            r["read_at"].strftime("%Y-%m-%dT%H:%M:%SZ"),
            r["temperature_c"],
            r["humidity_pct"],
            r["battery_v"],
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d06_coldstore_celsius.csv", rows)
    addd(
        "d06",
        "i2",
        "d06_coldstore_celsius.csv",
        recs,
        ["renamed", "unit"],
        "review",
        "the gateway now reports Celsius: 'Temp (F)' became 'Temp (C)'",
    )

    epoch = dt.datetime(1970, 1, 1)
    recs = sensor_records(rng, dt.datetime(2026, 4, 1, 8, 0), ["CS-01", "CS-02", "CS-03", "CS-04"])
    rows = [["device", "ts", "temp", "rh", "bat"]]
    rows += [
        [
            r["device_id"].lower(),
            int((r["read_at"] - epoch).total_seconds()),
            r["temperature_c"],
            r["humidity_pct"],
            r["battery_v"],
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d07_coldstore_dump_0401.csv", rows, delim=";")
    addd("d07", "i1", "d07_coldstore_dump_0401.csv", recs, ["read"], "loaded", "the delimiter changed to ';'")

    recs = crm_records(rng, start_id=8400)
    rows = [
        [
            "Record ID",
            "First Name",
            "Last Name",
            "E-mail",
            "Mobile",
            "Organisation",
            "Date Created",
            "Total Spend",
            "Notes",
        ]
    ]
    rows += [
        [
            r["contact_id"],
            r["first"],
            r["last"],
            r["email"] or "",
            nz_local(r["phone"], "local"),
            r["company"],
            r["created_on"].strftime("%d-%b-%Y"),
            usd(r["lifetime_value"]),
            "",
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d08_crm_contacts_split_name.csv", rows)
    addd(
        "d08",
        "c2",
        "d08_crm_contacts_split_name.csv",
        recs,
        ["split"],
        "review",
        "'Name' was split into 'First Name' and 'Last Name'",
    )

    recs = crm_records(rng, start_id=9600)
    rows = [["Contact Name", "Mail", "Phone Number", "Employer", "Created Date", "Revenue"]]
    rows += [
        [
            r["full_name"],
            r["email"] or "",
            nz_local(r["phone"], "bare"),
            r["company"],
            f"{r['created_on'].isoformat()}T10:22:00",
            eu_num(r["lifetime_value"]),
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d09_crm_semicolon_no_id.csv", rows, delim=";")
    addd(
        "d09",
        "c5",
        "d09_crm_semicolon_no_id.csv",
        recs,
        ["removed"],
        "blocked",
        "the 'ID' column (the table key) is missing",
    )

    recs = supplier_records(rng, "Nordwerk Tools GmbH", "EUR", dt.date(2026, 5, 1), "NW", units=("EA", "PK", "BOX"))
    rows = [["Artikel", "Bezeichnung", "Einheit", "VPE", "Preis", "Gueltig ab"]]
    rows += [
        [
            r["sku"],
            r["description"],
            r["unit"],
            r["pack_size"],
            f"EUR {r['price']:.2f}",
            r["valid_from"].strftime("%d.%m.%Y"),
        ]
        for r in recs
    ]
    write_csv(DRIFT / "d10_nordwerk_preisliste_mai.csv", rows, delim=";", pre=["Lieferant: Nordwerk Tools GmbH"])
    addd(
        "d10",
        "s3",
        "d10_nordwerk_preisliste_mai.csv",
        recs,
        ["format"],
        "loaded",
        "prices changed from '12,50 €' to 'EUR 12.50'",
    )


def extra_files() -> None:
    rng = random.Random(88)
    rows = [["Card", "Driver", "Station", "Litres", "Fuel Type", "Pumped At", "Cost"]]
    for _ in range(12):
        rows.append(
            [
                f"FC-{rng.randint(100, 999)}",
                rng.choice(FIRST),
                rng.choice(["Albany", "Takapuna", "Penrose"]),
                f"{rng.uniform(20, 80):.2f}",
                rng.choice(["91", "95", "Diesel"]),
                f"2026-03-{rng.randint(1, 28):02d} {rng.randint(6, 19):02d}:{rng.randint(0, 59):02d}",
                f"{rng.uniform(50, 220):.2f}",
            ]
        )
    write_csv(EXTRA / "fleet_fuel_cards.csv", rows)


def main() -> None:
    for d in (FILES, GOLD, DRIFT, EXTRA):
        shutil.rmtree(d, ignore_errors=True)
        d.mkdir(parents=True)
    files: list[Manifest] = []
    drift: list[Manifest] = []

    def add(
        fid: str,
        name: str,
        source: str,
        table: str,
        gold: list[dict[str, Any]],
        quirks: list[str],
        planted: list[dict[str, Any]] | None = None,
        quarantine: int = 0,
    ) -> None:
        write_gold(GOLD / f"{fid}.csv", table, gold)
        files.append(
            {
                "id": fid,
                "file": f"files/{name}",
                "source": source,
                "table": table,
                "rows": len(gold),
                "quirks": quirks,
                "planted": planted or [],
                "quarantine": quarantine,
                "gold": f"gold/{fid}.csv",
            }
        )

    for fn in (bank_files, supplier_files, pos_files, timesheet_files, sensor_files, crm_files):
        fn(add)
    by_id = {f["id"]: f for f in files}

    def addd(
        did: str, base: str, name: str, gold: list[dict[str, Any]], kinds: list[str], expect: str, what: str
    ) -> None:
        table = by_id[base]["table"]
        write_gold(GOLD / f"{did}.csv", table, gold)
        drift.append(
            {
                "id": did,
                "base": base,
                "file": f"drift/{name}",
                "table": table,
                "rows": len(gold),
                "kinds": kinds,
                "expect": expect,
                "what": what,
                "gold": f"gold/{did}.csv",
            }
        )

    drift_files(addd)
    extra_files()
    manifest = {"files": files, "drift": drift, "keys": KEYS}
    (HERE / "manifest.json").write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    print(f"{len(files)} files, {len(drift)} drift scenarios, {sum(len(f['planted']) for f in files)} planted problems")


if __name__ == "__main__":
    main()
