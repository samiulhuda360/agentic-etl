"""Turn a messy file into a table of raw text, and write down exactly how it was read.

Reading is deterministic and rule-based on purpose: the encoding, delimiter, header row, merged headers, sheets,
JSON record path and fixed-width boundaries are all detected here and recorded in a `ReadSpec`. Nothing is parsed
into numbers or dates yet; that is the transform's job. Every line that is not data (a title, a preamble, a
"Closing balance" footer, a cover sheet) is kept verbatim in `RawTable.skipped`, so it shows up in the run record
instead of disappearing.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import polars as pl

DELIMITERS = [",", ";", "\t", "|"]
FOOTER = re.compile(
    r"^\s*(total|totals|grand total|sub ?total|closing balance|opening balance|end of|records?:|row count|"
    r"\d+ records?)\b",
    re.I,
)
META = re.compile(r"^\s*([A-Za-z][A-Za-z /#._-]{1,40}?)\s*[:=]\s*(.+?)\s*$")


@dataclass
class ReadSpec:
    format: str  # csv | excel | json | ndjson | fixed
    encoding: str | None = None
    delimiter: str | None = None
    header_row: int = 0
    header_rows: int = 1
    sheets: list[str] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)
    record_path: str | None = None
    explode: str | None = None
    widths: list[list[int]] | None = None

    def to_dict(self) -> dict[str, Any]:
        return {k: v for k, v in asdict(self).items() if v not in (None, [], "")}


@dataclass
class RawTable:
    path: Path
    df: pl.DataFrame  # every column is String; meta values arrive as constant "@name" columns
    spec: ReadSpec
    meta: dict[str, str]
    skipped: list[str]
    sha256: str

    @property
    def data_columns(self) -> list[str]:
        return [c for c in self.df.columns if not c.startswith("@")]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read(path: Path) -> RawTable:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        df, spec, meta, skipped = _read_excel(path)
    else:
        raw = path.read_bytes()
        enc = detect_encoding(raw)
        text = raw.decode(enc)
        if text.startswith("﻿"):
            text = text[1:]
        stripped = text.lstrip()
        if suffix in (".json", ".ndjson", ".jsonl") or stripped[:1] in "{[":
            df, spec, meta, skipped = _read_json(text)
        else:
            df, spec, meta, skipped = _read_text(text)
        spec.encoding = enc
    for k, v in meta.items():
        df = df.with_columns(pl.lit(v, dtype=pl.String).alias(f"@{k}"))
    return RawTable(path, df, spec, meta, skipped, sha256(path))


# ---------------------------------------------------------------- encodings


def detect_encoding(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        return "cp1252"


# ---------------------------------------------------------------- delimited and fixed-width text


def _rows(text: str, delim: str) -> list[list[str]]:
    return [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]


def sniff_delimiter(text: str) -> tuple[str | None, int]:
    best: tuple[str | None, int, int] = (None, 1, 0)
    for d in DELIMITERS:
        rows = _rows(text, d)
        if not rows:
            continue
        modal, n = Counter(len(r) for r in rows).most_common(1)[0]
        if modal > 1 and n * modal > best[2]:
            best = (d, modal, n * modal)
    return best[0], best[1]


def _names(cells: list[Any]) -> list[str]:
    out: list[str] = []
    seen: Counter[str] = Counter()
    for i, c in enumerate(cells):
        name = re.sub(r"\s+", " ", str(c if c is not None else "")).strip() or f"column_{i + 1}"
        seen[name] += 1
        out.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return out


def _meta_from(cells: list[str]) -> tuple[str, str] | None:
    vals = [c.strip() for c in cells if c and c.strip()]
    if len(vals) == 1:
        m = META.match(vals[0])
        if m:
            return m.group(1).strip(), m.group(2).strip()
    if len(vals) == 2 and vals[0].rstrip(":").strip() and not re.search(r"\d", vals[0]):
        return vals[0].rstrip(":").strip(), vals[1]
    return None


def _split_footer(rows: list[list[str]], width: int) -> tuple[list[list[str]], list[list[str]]]:
    end = len(rows)
    while end > 0:
        r = rows[end - 1]
        cells = [c for c in r if c.strip()]
        if len(r) != width or (cells and FOOTER.match(cells[0])) or any(FOOTER.match(c) for c in r if c.strip()):
            end -= 1
            continue
        break
    return rows[:end], rows[end:]


def _looks_like_header(row: list[str]) -> bool:
    cells = [c.strip() for c in row if c.strip()]
    if not cells:
        return False
    numeric = sum(bool(re.fullmatch(r"[-+$€£(]?[\d.,:/ -]+[)%]?", c)) for c in cells)
    return numeric <= len(cells) // 4


def _read_text(text: str) -> tuple[pl.DataFrame, ReadSpec, dict[str, str], list[str]]:
    delim, width = sniff_delimiter(text)
    if delim is None or width < 3:
        return _read_fixed(text)
    rows = _rows(text, delim)
    start = 0
    for i, r in enumerate(rows):
        nxt = rows[i + 1 : i + 4]
        if len(r) == width and _looks_like_header(r) and nxt and all(len(x) == width for x in nxt[:1]):
            start = i
            break
    meta: dict[str, str] = {}
    skipped = []
    for r in rows[:start]:
        skipped.append(delim.join(r))
        kv = _meta_from(r)
        if kv:
            meta[kv[0]] = kv[1]
    header = _names(rows[start])
    body, footer = _split_footer(rows[start + 1 :], width)
    skipped += [delim.join(r) for r in footer]
    good = []
    for r in body:
        if len(r) == width:
            good.append(r)
        else:
            skipped.append(delim.join(r))
    df = pl.DataFrame({h: [r[j] for r in good] for j, h in enumerate(header)}, schema={h: pl.String for h in header})
    spec = ReadSpec(format="csv", delimiter=delim, header_row=start)
    return df, spec, meta, skipped


def _fixed_fields(lines: list[str]) -> list[tuple[int, int]]:
    width = max(len(x) for x in lines)
    blank = [all(i >= len(x) or x[i] == " " for x in lines) for i in range(width)]
    fields: list[tuple[int, int]] = []
    i = 0
    while i < width:
        if blank[i]:
            i += 1
            continue
        j = i
        while j < width and not blank[j]:
            j += 1
        fields.append((i, j))
        i = j
    # A field must hold the start of a header token; otherwise it is the tail of the previous field.
    header = lines[0]
    starts = {m.start() for m in re.finditer(r"\S+", header)}
    merged: list[tuple[int, int]] = []
    for a, b in fields:
        if merged and not any(a <= s < b for s in starts):
            merged[-1] = (merged[-1][0], b)
        else:
            merged.append((a, b))
    return merged


def _read_fixed(text: str) -> tuple[pl.DataFrame, ReadSpec, dict[str, str], list[str]]:
    lines = [ln.rstrip("\r") for ln in text.split("\n") if ln.strip()]
    footer: list[str] = []
    while lines and FOOTER.match(lines[-1].strip()):
        footer.insert(0, lines.pop())
    start, fields = 0, _fixed_fields(lines)
    for s in range(min(6, len(lines) - 1)):
        f = _fixed_fields(lines[s:])
        if len(f) >= 3 and len(lines[s].split()) == len(f):
            start, fields = s, f
            break
    meta: dict[str, str] = {}
    skipped = []
    for ln in lines[:start]:
        skipped.append(ln)
        kv = _meta_from([ln])
        if kv:
            meta[kv[0]] = kv[1]
    block = lines[start:]
    header = _names([block[0][a:b].strip() for a, b in fields])
    data = []
    for ln in block[1:]:
        data.append([ln[a:b].strip() if a < len(ln) else "" for a, b in fields])
    df = pl.DataFrame({h: [r[j] for r in data] for j, h in enumerate(header)}, schema={h: pl.String for h in header})
    spec = ReadSpec(format="fixed", header_row=start, widths=[[a, b] for a, b in fields])
    return df, spec, meta, skipped + footer


# ---------------------------------------------------------------- Excel


def _cell_text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == dt.time(0) else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v)


def _sheet_table(ws: Any) -> tuple[list[str], list[list[str | None]], dict[str, str], list[str], int] | None:
    grid: list[list[Any]] = [list(r) for r in ws.iter_rows(values_only=True)]
    if not grid:
        return None
    merged_in_row: dict[int, int] = Counter()
    for rng in ws.merged_cells.ranges:
        top = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(grid) else None
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                    grid[r - 1][c - 1] = top
        if rng.max_col > rng.min_col:
            merged_in_row[rng.min_row - 1] += 1
    text = [[_cell_text(v) for v in row] for row in grid]
    counts = [sum(1 for v in row if v not in (None, "")) for row in text]
    width = Counter(c for c in counts if c > 1).most_common(1)
    if not width:
        return None
    w = width[0][0]
    hdr = next(
        (
            i
            for i, c in enumerate(counts)
            if c >= max(2, int(0.6 * w)) and _looks_like_header([v or "" for v in text[i]])
        ),
        None,
    )
    if hdr is None:
        return None
    two_rows = (
        merged_in_row.get(hdr, 0) > 0 and hdr + 1 < len(text) and _looks_like_header([v or "" for v in text[hdr + 1]])
    )
    meta: dict[str, str] = {}
    skipped = []
    for row in text[:hdr]:
        cells = [v for v in row if v]
        if cells:
            skipped.append(" | ".join(cells))
            kv = _meta_from(cells)
            if kv:
                meta[kv[0]] = kv[1]
    if two_rows:
        top, bottom = text[hdr], text[hdr + 1]
        names = []
        for a, b in zip(top, bottom, strict=False):
            if a and b and a != b:
                names.append(f"{a} {b}")
            else:
                names.append(b or a or "")
        first = hdr + 2
    else:
        names = [v or "" for v in text[hdr]]
        first = hdr + 1
    keep = [i for i, n in enumerate(names) if n or any(r[i] not in (None, "") for r in text[first:] if i < len(r))]
    header = _names([names[i] for i in keep])
    body = [[r[i] if i < len(r) else None for i in keep] for r in text[first:]]
    body = [r for r in body if any(v not in (None, "") for v in r)]
    end = len(body)
    while end > 0 and any(v and FOOTER.match(v) for v in body[end - 1]):
        end -= 1
    skipped += [" | ".join(v for v in r if v) for r in body[end:]]
    return header, body[:end], meta, skipped, 2 if two_rows else 1


def _read_excel(path: Path) -> tuple[pl.DataFrame, ReadSpec, dict[str, str], list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    tables = {}
    skipped: list[str] = []
    for ws in wb.worksheets:
        t = _sheet_table(ws)
        if t is None or not t[1]:
            skipped.append(f"sheet '{ws.title}' has no data table")
            continue
        tables[ws.title] = t
    if not tables:
        raise ValueError(f"{path.name}: no sheet holds a data table")
    sig = Counter(tuple(t[0]) for t in tables.values()).most_common(1)[0][0]
    use = [s for s, t in tables.items() if tuple(t[0]) == sig]
    other = [s for s in tables if s not in use]
    for s in other:
        skipped.append(f"sheet '{s}' has a different layout ({len(tables[s][1])} rows), not stacked")
    header = list(sig)
    rows: list[list[str | None]] = []
    sheet_col: list[str] = []
    meta: dict[str, str] = {}
    for s in use:
        h, body, m, sk, _ = tables[s]
        rows += body
        sheet_col += [s] * len(body)
        meta.update(m)
        skipped += [f"[{s}] {x}" for x in sk]
    df = pl.DataFrame({h: [r[j] for r in rows] for j, h in enumerate(header)}, schema={h: pl.String for h in header})
    if len(use) > 1:
        df = df.with_columns(pl.Series("@sheet", sheet_col, dtype=pl.String))
    spec = ReadSpec(
        format="excel",
        sheets=use,
        skipped_sheets=[w.title for w in wb.worksheets if w.title not in use],
        header_rows=tables[use[0]][4],
    )
    return df, spec, meta, skipped


# ---------------------------------------------------------------- JSON


def _scalar(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        return repr(v)
    if isinstance(v, (dict, list)):
        return json.dumps(v, sort_keys=True, ensure_ascii=False)
    return str(v)


def _flatten(d: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    for k, v in d.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(_flatten(v, key + "."))
        else:
            out[key] = v
    return out


def _find_records(obj: Any, path: str = "") -> tuple[str, list[dict[str, Any]]] | None:
    best: tuple[str, list[dict[str, Any]]] | None = None
    if isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj):
        return path or "$", obj
    if isinstance(obj, dict):
        for k, v in obj.items():
            found = _find_records(v, f"{path}.{k}" if path else k)
            if found and (best is None or len(found[1]) > len(best[1])):
                best = found
    return best


def _read_json(text: str) -> tuple[pl.DataFrame, ReadSpec, dict[str, str], list[str]]:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    meta: dict[str, str] = {}
    if len(lines) > 1 and all(ln.lstrip().startswith("{") for ln in lines):
        try:
            records = [json.loads(ln) for ln in lines]
            spec = ReadSpec(format="ndjson", record_path="$")
        except json.JSONDecodeError:
            records = []
    else:
        records = []
    if not records:
        obj = json.loads(text)
        found = _find_records(obj)
        if not found:
            raise ValueError("no list of records found in the JSON")
        rpath, records = found
        spec = ReadSpec(format="json", record_path=rpath)
        node: Any = obj
        parts = [] if rpath == "$" else rpath.split(".")
        for depth, part in enumerate(parts):
            if not isinstance(node, dict):
                break
            prefix = ".".join(parts[:depth])
            flat = _flatten({k: v for k, v in node.items() if k != part}, prefix + "." if prefix else "")
            meta.update({k: s for k, v in flat.items() if not isinstance(v, list) and (s := _scalar(v)) is not None})
            node = node[part]
    flat_rows = [_flatten(r) for r in records]
    child_key = None
    sizes: Counter[str] = Counter()
    for r in flat_rows:
        for k, v in r.items():
            if isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
                sizes[k] += len(v)
    if sizes:
        child_key = sizes.most_common(1)[0][0]
        spec.explode = child_key
    rows: list[dict[str, Any]] = []
    for r in flat_rows:
        children = r.pop(child_key, None) if child_key else None
        if children:
            for i, ch in enumerate(children, 1):
                row = dict(r)
                row[f"{child_key}.#"] = i
                for k, v in _flatten(ch).items():
                    row[f"{child_key}.{k}"] = v
                rows.append(row)
        else:
            rows.append(r)
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    df = pl.DataFrame({c: [_scalar(r.get(c)) for r in rows] for c in cols}, schema={c: pl.String for c in cols})
    return df, spec, meta, []
